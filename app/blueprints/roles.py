"""
Manajemen Role & matrix hak akses (Create/Read/Edit/Delete per
object) - khusus administrator. Termasuk export/import matrix
dalam bentuk Excel.
"""

import io

from datetime import datetime

from flask import (
    render_template,
    request,
    redirect,
    url_for,
    flash,
    send_file,
)
from flask_babel import gettext as _

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

from .. import db
from ..models import (
    Role,
    Permission,
    User,
    PERMISSION_OBJECTS,
    PERMISSION_OBJECT_KEYS,
    PERMISSION_ACTIONS,
    ROLE_ADMINISTRATOR,
)
from ..auth import administrator_required

from . import main
from .helpers import log_action


@main.route("/roles")
@administrator_required
def roles():

    items = (
        Role.query
        .order_by(Role.name.asc())
        .all()
    )

    user_counts = {
        row[0]: row[1]
        for row in (
            db.session.query(
                User.role,
                db.func.count(User.id),
            )
            .group_by(User.role)
            .all()
        )
    }

    return render_template(
        "roles/list.html",
        roles=items,
        user_counts=user_counts,
    )


@main.route(
    "/roles/add",
    methods=["GET", "POST"],
)
@administrator_required
def add_role():

    if request.method == "POST":

        name = request.form.get(
            "name",
            "",
        ).strip()

        description = request.form.get(
            "description",
            "",
        ).strip()

        if not name:

            flash(
                _("Role name is required."),
                "danger",
            )

            return redirect(
                url_for("main.add_role")
            )

        if Role.query.filter_by(name=name).first():

            flash(
                _("Role name is already in use."),
                "danger",
            )

            return redirect(
                url_for("main.add_role")
            )

        role = Role(
            name=name,
            description=description,
        )

        db.session.add(role)
        db.session.flush()

        # Buat baris permission kosong (semua False) untuk tiap
        # object, supaya langsung muncul di halaman matrix.
        for object_name in PERMISSION_OBJECT_KEYS:

            db.session.add(
                Permission(
                    role_id=role.id,
                    object_name=object_name,
                )
            )

        db.session.commit()

        log_action(
            "create",
            "Role",
            entity_label=role.name,
        )

        flash(
            _("Role created successfully. Set its permissions below."),
            "success",
        )

        return redirect(
            url_for("main.role_permissions", role_id=role.id)
        )

    return render_template("roles/add.html")


@main.route(
    "/roles/<int:role_id>/edit",
    methods=["GET", "POST"],
)
@administrator_required
def edit_role(role_id):

    role = Role.query.get_or_404(role_id)

    if request.method == "POST":

        if role.is_system:

            flash(
                _("Built-in roles cannot be renamed."),
                "danger",
            )

            return redirect(
                url_for("main.roles")
            )

        name = request.form.get(
            "name",
            "",
        ).strip()

        description = request.form.get(
            "description",
            "",
        ).strip()

        if not name:

            flash(
                _("Role name is required."),
                "danger",
            )

            return redirect(
                url_for("main.edit_role", role_id=role.id)
            )

        duplicate = Role.query.filter(
            Role.name == name,
            Role.id != role.id,
        ).first()

        if duplicate:

            flash(
                _("Role name is already in use."),
                "danger",
            )

            return redirect(
                url_for("main.edit_role", role_id=role.id)
            )

        old_name = role.name

        # Kalau nama role berubah, samakan juga User.role yang
        # sedang memakai nama role lama, supaya tidak ada user
        # yang "kehilangan" role-nya.
        if old_name != name:

            User.query.filter_by(role=old_name).update(
                {"role": name}
            )

        role.name = name
        role.description = description

        db.session.commit()

        log_action(
            "update",
            "Role",
            entity_label=role.name,
        )

        flash(
            _("Role updated successfully."),
            "success",
        )

        return redirect(
            url_for("main.roles")
        )

    return render_template(
        "roles/edit.html",
        role=role,
    )


@main.route(
    "/roles/<int:role_id>/permissions",
    methods=["GET", "POST"],
)
@administrator_required
def role_permissions(role_id):

    role = Role.query.get_or_404(role_id)

    if request.method == "POST":

        for object_name in PERMISSION_OBJECT_KEYS:

            permission = role.permission_for(object_name)

            if not permission:

                permission = Permission(
                    role_id=role.id,
                    object_name=object_name,
                )

                db.session.add(permission)

            for action in PERMISSION_ACTIONS:

                field_name = f"perm_{object_name}_{action}"

                setattr(
                    permission,
                    f"can_{action}",
                    field_name in request.form,
                )

        db.session.commit()

        log_action(
            "update",
            "Role",
            entity_label=role.name,
            description="Ubah hak akses (permission matrix)",
        )

        flash(
            _("Role permissions saved successfully."),
            "success",
        )

        return redirect(
            url_for("main.role_permissions", role_id=role.id)
        )

    permissions_by_object = {
        object_name: role.permission_for(object_name)
        for object_name in PERMISSION_OBJECT_KEYS
    }

    return render_template(
        "roles/permissions.html",
        role=role,
        permission_objects=PERMISSION_OBJECTS,
        permission_actions=PERMISSION_ACTIONS,
        permissions_by_object=permissions_by_object,
    )


@main.route(
    "/roles/<int:role_id>/delete",
    methods=["POST"],
)
@administrator_required
def delete_role(role_id):

    role = Role.query.get_or_404(role_id)

    if role.is_system:

        flash(
            _("Built-in roles cannot be deleted."),
            "danger",
        )

        return redirect(
            url_for("main.roles")
        )

    users_with_role = User.query.filter_by(
        role=role.name
    ).count()

    if users_with_role > 0:

        flash(
            _(
                "This role is still used by %(count)s user(s). "
                "Move those users to another role before deleting it."
            )
            % {"count": users_with_role},
            "danger",
        )

        return redirect(
            url_for("main.roles")
        )

    deleted_name = role.name

    db.session.delete(role)
    db.session.commit()

    log_action(
        "delete",
        "Role",
        entity_label=deleted_name,
    )

    flash(
        _("Role deleted successfully."),
        "success",
    )

    return redirect(
        url_for("main.roles")
    )


# ============================================================
# EXPORT / IMPORT MATRIX HAK AKSES
# ============================================================


@main.route("/roles/export")
@administrator_required
def export_roles():

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Roles & Permissions"

    headers = (
        ["Role", "Deskripsi"]
        + [
            f"{label} - {action.capitalize()}"
            for _key, label in PERMISSION_OBJECTS
            for action in PERMISSION_ACTIONS
        ]
    )

    for col, header in enumerate(headers, start=1):

        cell = worksheet.cell(row=1, column=col, value=header)

        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="343A40",
            end_color="343A40",
            fill_type="solid",
        )
        cell.alignment = Alignment(
            horizontal="center",
            wrap_text=True,
        )

    row_num = 2

    for role in Role.query.order_by(Role.name.asc()).all():

        worksheet.cell(row=row_num, column=1, value=role.name)
        worksheet.cell(row=row_num, column=2, value=role.description or "")

        col_num = 3

        for object_name, _label in PERMISSION_OBJECTS:

            permission = role.permission_for(object_name)

            for action in PERMISSION_ACTIONS:

                value = (
                    "Y"
                    if permission and permission.allows(action)
                    else "N"
                )

                worksheet.cell(row=row_num, column=col_num, value=value)

                col_num += 1

        row_num += 1

    for col_num in range(1, len(headers) + 1):

        worksheet.column_dimensions[
            worksheet.cell(row=1, column=col_num).column_letter
        ].width = 14

    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 30

    buffer = io.BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"roles_permissions_{timestamp}.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


@main.route(
    "/roles/import",
    methods=["GET", "POST"],
)
@administrator_required
def import_roles():

    if request.method == "POST":

        upload = request.files.get("file")

        if not upload or not upload.filename:

            flash(
                _("Please select the exported Excel file first."),
                "danger",
            )

            return redirect(
                url_for("main.import_roles")
            )

        try:

            workbook = load_workbook(upload)
            worksheet = workbook.active

        except Exception:

            flash(
                _(
                    "The file could not be read. Make sure it is an "
                    ".xlsx file exported from this page."
                ),
                "danger",
            )

            return redirect(
                url_for("main.import_roles")
            )

        header_row = next(worksheet.iter_rows(min_row=1, max_row=1))

        headers = [
            (cell.value or "").strip()
            for cell in header_row
        ]

        expected_columns = [
            f"{label} - {action.capitalize()}"
            for _key, label in PERMISSION_OBJECTS
            for action in PERMISSION_ACTIONS
        ]

        if headers[:2] != ["Role", "Deskripsi"] or headers[2:] != expected_columns:

            flash(
                _(
                    "The file's column format is invalid. Use a file "
                    "exported from this page as the template."
                ),
                "danger",
            )

            return redirect(
                url_for("main.import_roles")
            )

        role_count = 0
        permission_count = 0

        for row in worksheet.iter_rows(min_row=2, values_only=True):

            if not row or not row[0]:
                continue

            role_name = str(row[0]).strip()
            description = str(row[1] or "").strip()

            role = Role.query.filter_by(name=role_name).first()

            if not role:

                role = Role(
                    name=role_name,
                    description=description,
                )

                db.session.add(role)
                db.session.flush()

                role_count += 1

            elif not role.is_system:

                role.description = description

            col_index = 2

            for object_name, _label in PERMISSION_OBJECTS:

                permission = role.permission_for(object_name)

                if not permission:

                    permission = Permission(
                        role_id=role.id,
                        object_name=object_name,
                    )

                    db.session.add(permission)

                for action in PERMISSION_ACTIONS:

                    cell_value = row[col_index] if col_index < len(row) else None

                    setattr(
                        permission,
                        f"can_{action}",
                        str(cell_value).strip().upper() == "Y",
                    )

                    col_index += 1

                permission_count += 1

        db.session.commit()

        log_action(
            "update",
            "Role",
            description=(
                f"Import matrix hak akses "
                f"({role_count} role baru, {permission_count} baris izin)"
            ),
        )

        flash(
            _(
                "Import successful. %(role_count)s new role(s) created, "
                "permissions for %(permission_count)s object row(s) updated."
            )
            % {
                "role_count": role_count,
                "permission_count": permission_count,
            },
            "success",
        )

        return redirect(
            url_for("main.roles")
        )

    return render_template("roles/import.html")
