"""
Warranty management: monitoring, alert, export, riwayat perubahan.
"""

import io

from datetime import datetime, date

from flask import render_template, request, send_file

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from .. import db
from ..models import Asset, Category, Location, Department, Vendor, WarrantyHistory
from ..auth import requires_permission

from . import main
from .helpers import get_warranty_status

@main.route("/warranty")
@requires_permission("warranty", "read")
def warranty_management():

    # --------------------------------------------------------
    # FILTER
    # --------------------------------------------------------

    selected_status = request.args.get(
        "status",
        ""
    ).strip()

    search = request.args.get(
        "search",
        ""
    ).strip()

    # --------------------------------------------------------
    # AMBIL SEMUA ASSET
    # --------------------------------------------------------

    assets_query = (
        Asset.query
        .order_by(
            Asset.asset_name.asc(),
            Asset.id.asc()
        )
    )

    # --------------------------------------------------------
    # SEARCH
    # --------------------------------------------------------

    if search:

        search_pattern = f"%{search}%"

        assets_query = assets_query.filter(
            db.or_(
                Asset.asset_tag.ilike(
                    search_pattern
                ),
                Asset.asset_name.ilike(
                    search_pattern
                ),
                Asset.brand.ilike(
                    search_pattern
                ),
                Asset.model.ilike(
                    search_pattern
                ),
                Asset.serial_number.ilike(
                    search_pattern
                ),
                Asset.vendor.ilike(
                    search_pattern
                )
            )
        )

    assets = assets_query.all()

    # --------------------------------------------------------
    # WARRANTY DATA
    # --------------------------------------------------------

    today = date.today()

    warranty_records = []

    for asset in assets:

        warranty_info = get_warranty_status(
            asset.purchase_date,
            asset.warranty,
            today=today
        )

        record = {
            "asset": asset,
            "end_date": warranty_info.get(
                "end_date"
            ),
            "days_left": warranty_info.get(
                "days_left"
            ),
            "status": warranty_info.get(
                "status"
            )
        }

        warranty_records.append(
            record
        )

    # --------------------------------------------------------
    # FILTER STATUS
    # --------------------------------------------------------

    if selected_status:

        warranty_records = [
            item
            for item in warranty_records
            if item["status"]
            == selected_status
        ]

    # --------------------------------------------------------
    # SORT
    # --------------------------------------------------------

    def warranty_sort_key(item):

        end_date = item.get(
            "end_date"
        )

        if end_date:
            return end_date

        return date.max

    warranty_records.sort(
        key=warranty_sort_key
    )

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    all_warranty_records = []

    for asset in assets:

        warranty_info = get_warranty_status(
            asset.purchase_date,
            asset.warranty,
            today=today
        )

        all_warranty_records.append(
            warranty_info
        )

    warranty_active_count = sum(
        1
        for item in all_warranty_records
        if item.get("status")
        == "active"
    )

    warranty_expiring_count = sum(
        1
        for item in all_warranty_records
        if item.get("status")
        == "expiring"
    )

    warranty_expired_count = sum(
        1
        for item in all_warranty_records
        if item.get("status")
        == "expired"
    )

    warranty_unknown_count = sum(
        1
        for item in all_warranty_records
        if item.get("status")
        not in (
            "active",
            "expiring",
            "expired"
        )
    )

    # --------------------------------------------------------
    # RENDER
    # --------------------------------------------------------

    return render_template(
        "warranty/index.html",

        warranty_records=warranty_records,

        warranty_active_count=(
            warranty_active_count
        ),

        warranty_expiring_count=(
            warranty_expiring_count
        ),

        warranty_expired_count=(
            warranty_expired_count
        ),

        warranty_unknown_count=(
            warranty_unknown_count
        ),

        selected_status=(
            selected_status
        ),

        search=search
    )

# ============================================================
# WARRANTY ALERT
# ============================================================


@main.route("/warranty/alerts")
@requires_permission("warranty", "read")
def warranty_alerts():

    today = date.today()

    assets = (
        Asset.query
        .order_by(
            Asset.asset_name.asc(),
            Asset.id.asc()
        )
        .all()
    )

    expired_records = []
    urgent_records = []
    warning_records = []
    active_records = []
    unknown_records = []

    for asset in assets:

        warranty_info = get_warranty_status(
            asset.purchase_date,
            asset.warranty,
            today=today
        )

        status = warranty_info.get(
            "status",
            "unknown"
        )

        days_left = warranty_info.get(
            "days_left"
        )

        end_date = warranty_info.get(
            "end_date"
        )

        record = {
            "asset": asset,
            "status": status,
            "days_left": days_left,
            "end_date": end_date,
        }

        # ----------------------------------------------------
        # EXPIRED
        # ----------------------------------------------------

        if status == "expired":

            expired_records.append(
                record
            )

        # ----------------------------------------------------
        # URGENT <= 7 DAYS
        # ----------------------------------------------------

        elif (
            status == "expiring"
            and days_left is not None
            and days_left <= 7
        ):

            urgent_records.append(
                record
            )

        # ----------------------------------------------------
        # WARNING <= 30 DAYS
        # ----------------------------------------------------

        elif (
            status == "expiring"
            and days_left is not None
            and days_left > 7
        ):

            warning_records.append(
                record
            )

        # ----------------------------------------------------
        # ACTIVE
        # ----------------------------------------------------

        elif status == "active":

            active_records.append(
                record
            )

        # ----------------------------------------------------
        # UNKNOWN
        # ----------------------------------------------------

        else:

            unknown_records.append(
                record
            )

    # --------------------------------------------------------
    # SORT
    # --------------------------------------------------------

    expired_records.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 0
        )
    )

    urgent_records.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 999999
        )
    )

    warning_records.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 999999
        )
    )

    active_records.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 999999
        )
    )

    return render_template(
        "warranty/alerts.html",

        expired_records=expired_records,

        urgent_records=urgent_records,

        warning_records=warning_records,

        active_records=active_records,

        unknown_records=unknown_records,

        expired_count=len(
            expired_records
        ),

        urgent_count=len(
            urgent_records
        ),

        warning_count=len(
            warning_records
        ),

        active_count=len(
            active_records
        ),

        unknown_count=len(
            unknown_records
        ),

        today=today
    )

# ============================================================
# EXPORT WARRANTY DETAIL
# ============================================================


@main.route("/warranty/export")
@requires_permission("warranty", "read")
def export_warranty():

    from openpyxl import Workbook
    from openpyxl.styles import (
        Font,
        Alignment,
        PatternFill,
    )

    # ========================================================
    # FILTER
    # ========================================================

    selected_status = (
        request.args.get(
            "status",
            "all"
        )
        .strip()
        .lower()
    )

    search = (
        request.args.get(
            "search",
            ""
        )
        .strip()
        .lower()
    )

    today = date.today()

    # ========================================================
    # AMBIL ASSET
    # ========================================================

    assets = (
        Asset.query
        .order_by(
            Asset.asset_name.asc(),
            Asset.id.asc()
        )
        .all()
    )

    # ========================================================
    # WARRANTY RECORDS
    # ========================================================

    warranty_records = []

    for asset in assets:

        # ----------------------------------------------------
        # SEARCH
        # ----------------------------------------------------

        if search:

            searchable_text = " ".join([
                str(asset.asset_tag or ""),
                str(asset.asset_name or ""),
                str(asset.brand or ""),
                str(asset.model or ""),
                str(asset.serial_number or ""),
                str(asset.vendor or ""),
                str(asset.location or ""),
                str(asset.department or ""),
                str(asset.pic or ""),
            ]).lower()

            if search not in searchable_text:
                continue

        # ----------------------------------------------------
        # WARRANTY STATUS
        # ----------------------------------------------------

        warranty_info = get_warranty_status(
            asset.purchase_date,
            asset.warranty,
            today=today
        )

        status = warranty_info.get(
            "status",
            "unknown"
        )

        days_left = warranty_info.get(
            "days_left"
        )

        end_date = warranty_info.get(
            "end_date"
        )

        # ----------------------------------------------------
        # FILTER STATUS
        # ----------------------------------------------------

        include_record = False

        if selected_status in (
            "",
            "all",
        ):

            include_record = True

        elif selected_status == "urgent":

            include_record = (
                status == "expiring"
                and days_left is not None
                and days_left <= 7
            )

        elif selected_status == "warning":

            include_record = (
                status == "expiring"
                and days_left is not None
                and days_left > 7
                and days_left <= 30
            )

        elif selected_status == "expiring":

            include_record = (
                status == "expiring"
            )

        elif selected_status == "expired":

            include_record = (
                status == "expired"
            )

        elif selected_status == "active":

            include_record = (
                status == "active"
            )

        elif selected_status == "unknown":

            include_record = (
                status == "unknown"
            )

        if not include_record:
            continue

        warranty_records.append({
            "asset": asset,
            "status": status,
            "days_left": days_left,
            "end_date": end_date,
        })

    # ========================================================
    # SORT
    # ========================================================

    if selected_status == "expired":

        warranty_records.sort(
            key=lambda item: (
                item["end_date"]
                if item["end_date"]
                else date.min
            ),
            reverse=True
        )

    else:

        warranty_records.sort(
            key=lambda item: (
                item["days_left"]
                if item["days_left"] is not None
                else 999999
            )
        )

    # ========================================================
    # WORKBOOK
    # ========================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Warranty Alert"

    # ========================================================
    # TITLE
    # ========================================================

    worksheet["A1"] = (
        "WARRANTY ALERT DETAIL - ASSET TRACKER"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=14,
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 26

    # ========================================================
    # INFORMATION
    # ========================================================

    filter_label = {
        "all": "Semua Warranty",
        "urgent": "Segera Habis ≤ 7 Hari",
        "warning": "Akan Habis 8–30 Hari",
        "expiring": "Akan Habis ≤ 30 Hari",
        "expired": "Warranty Expired",
        "active": "Warranty Aktif",
        "unknown": "Tidak Lengkap",
    }.get(
        selected_status,
        "Semua Warranty",
    )

    worksheet["A2"] = (
        "Tanggal Export: "
        + today.strftime("%d-%m-%Y")
        + " | Filter: "
        + filter_label
    )

    if search:

        worksheet["A2"] = (
            worksheet["A2"].value
            + " | Search: "
            + search
        )

    worksheet["A2"].alignment = Alignment(
        horizontal="center"
    )

    # ========================================================
    # HEADER
    # ========================================================

    headers = [
        "No",
        "Asset Tag",
        "Asset Name",
        "Category",
        "Brand",
        "Model",
        "Serial Number",
        "Location",
        "Department",
        "Vendor",
        "Warranty",
        "Tanggal Pembelian",
        "Tanggal Berakhir",
        "Sisa Hari",
        "Status",
    ]

    # 15 kolom = A:O
    worksheet.merge_cells(
        "A1:O1"
    )

    worksheet.merge_cells(
        "A2:O2"
    )

    for column_number, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=4,
            column=column_number,
        )

        cell.value = header

        cell.font = Font(
            bold=True,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # ========================================================
    # HEADER FILL
    # ========================================================

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for cell in worksheet[4]:

        cell.fill = header_fill

        cell.font = Font(
            bold=True,
        )

    # ========================================================
    # DATA
    # ========================================================

    for number, item in enumerate(
        warranty_records,
        start=1,
    ):

        asset = item["asset"]

        status = item["status"]

        days_left = item["days_left"]

        end_date = item["end_date"]

        # ----------------------------------------------------
        # STATUS LABEL
        # ----------------------------------------------------

        if status == "expired":

            status_label = "Expired"

        elif status == "expiring":

            if (
                days_left is not None
                and days_left <= 7
            ):

                status_label = (
                    "Segera Habis"
                )

            else:

                status_label = (
                    "Akan Habis"
                )

        elif status == "active":

            status_label = "Aktif"

        else:

            status_label = (
                "Tidak Lengkap"
            )

        # ----------------------------------------------------
        # ROW
        # ----------------------------------------------------

        worksheet.append([

            number,

            asset.asset_tag or "",

            asset.asset_name or "",

            asset.category or "",

            asset.brand or "",

            asset.model or "",

            asset.serial_number or "",

            asset.location or "",

            asset.department or "",

            asset.vendor or "",

            asset.warranty or "",

            (
                asset.purchase_date
                if asset.purchase_date
                else ""
            ),

            (
                end_date
                if end_date
                else ""
            ),

            (
                days_left
                if days_left is not None
                else ""
            ),

            status_label,

        ])

    # ========================================================
    # DATE FORMAT
    # ========================================================

    for row in worksheet.iter_rows(
        min_row=5,
        max_row=worksheet.max_row,
    ):

        # Tanggal Pembelian = L
        if row[11].value:

            row[11].number_format = (
                "dd-mm-yyyy"
            )

        # Tanggal Berakhir = M
        if row[12].value:

            row[12].number_format = (
                "dd-mm-yyyy"
            )

    # ========================================================
    # STATUS COLOR
    # ========================================================

    status_colors = {

        "Aktif":
            "C6EFCE",

        "Segera Habis":
            "FFC7CE",

        "Akan Habis":
            "FFEB9C",

        "Expired":
            "D9D9D9",

        "Tidak Lengkap":
            "D9E1F2",

    }

    for row in worksheet.iter_rows(
        min_row=5,
        max_row=worksheet.max_row,
    ):

        status_cell = row[14]

        status_name = (
            status_cell.value
        )

        if status_name in status_colors:

            status_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=status_colors[
                    status_name
                ],
            )

            status_cell.font = Font(
                bold=True,
            )

            status_cell.alignment = Alignment(
                horizontal="center",
            )

    # ========================================================
    # ALIGNMENT
    # ========================================================

    for row in worksheet.iter_rows(
        min_row=4,
        max_row=worksheet.max_row,
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="center",
            )

    # ========================================================
    # CENTER COLUMNS
    # ========================================================

    for row in worksheet.iter_rows(
        min_row=5,
        max_row=worksheet.max_row,
    ):

        for index in (
            0,
            11,
            12,
            13,
            14,
        ):

            row[index].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # ========================================================
    # FREEZE
    # ========================================================

    worksheet.freeze_panes = "A5"

    # ========================================================
    # AUTO FILTER
    # ========================================================

    worksheet.auto_filter.ref = (
        f"A4:O{worksheet.max_row}"
    )

    # ========================================================
    # COLUMN WIDTH
    # ========================================================

    column_widths = {

        "A": 8,
        "B": 18,
        "C": 30,
        "D": 20,
        "E": 18,
        "F": 22,
        "G": 24,
        "H": 20,
        "I": 20,
        "J": 20,
        "K": 18,
        "L": 18,
        "M": 18,
        "N": 14,
        "O": 18,

    }

    for column, width in column_widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # ========================================================
    # OUTPUT
    # ========================================================

    output = io.BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    filename = (
        "warranty_alert_"
        + datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )
        + ".xlsx"
    )

    return send_file(

        output,

        as_attachment=True,

        download_name=filename,

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

    )

# ============================================================
# WARRANTY ALERT DETAIL
# ============================================================


@main.route("/warranty/detail")
@requires_permission("warranty", "read")
def warranty_alert_detail():

    # ========================================================
    # FILTER
    # ========================================================

    selected_status = (
        request.args.get(
            "status",
            "all"
        )
        .strip()
        .lower()
    )

    search = (
        request.args.get(
            "search",
            ""
        )
        .strip()
    )

    # ========================================================
    # DEFAULT DATA
    # ========================================================

    warranty_active = []

    warranty_expiring = []

    warranty_expired = []

    warranty_unknown = []

    today = date.today()

    # ========================================================
    # AMBIL SEMUA ASSET
    # ========================================================

    all_assets = (
        Asset.query
        .order_by(
            Asset.asset_name.asc(),
            Asset.id.asc()
        )
        .all()
    )

    # ========================================================
    # PROSES WARRANTY
    # ========================================================

    for asset in all_assets:

        # ----------------------------------------------------
        # SEARCH
        # ----------------------------------------------------

        if search:

            searchable_text = " ".join([
                str(asset.asset_tag or ""),
                str(asset.asset_name or ""),
                str(asset.brand or ""),
                str(asset.model or ""),
                str(asset.serial_number or ""),
                str(asset.vendor or ""),
                str(asset.location or ""),
                str(asset.department or ""),
                str(asset.pic or ""),
            ]).lower()

            if search.lower() not in searchable_text:
                continue

        # ----------------------------------------------------
        # GET WARRANTY STATUS
        # ----------------------------------------------------

        warranty_info = get_warranty_status(
            asset.purchase_date,
            asset.warranty,
            today=today
        )

        status = warranty_info.get(
            "status",
            "unknown"
        )

        days_left = warranty_info.get(
            "days_left"
        )

        end_date = warranty_info.get(
            "end_date"
        )

        warranty_record = {
            "asset": asset,
            "status": status,
            "days_left": days_left,
            "end_date": end_date,
        }

        # ----------------------------------------------------
        # ACTIVE
        # ----------------------------------------------------

        if status == "active":

            warranty_active.append(
                warranty_record
            )

        # ----------------------------------------------------
        # EXPIRING
        # ----------------------------------------------------

        elif status == "expiring":

            warranty_expiring.append(
                warranty_record
            )

        # ----------------------------------------------------
        # EXPIRED
        # ----------------------------------------------------

        elif status == "expired":

            warranty_expired.append(
                warranty_record
            )

        # ----------------------------------------------------
        # UNKNOWN
        # ----------------------------------------------------

        else:

            warranty_unknown.append(
                warranty_record
            )

    # ========================================================
    # SORT
    # ========================================================

    warranty_expiring.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 999999
        )
    )

    warranty_expired.sort(
        key=lambda item: (
            item["end_date"]
            if item["end_date"] is not None
            else date.min
        ),
        reverse=True
    )

    warranty_active.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 999999
        )
    )

    warranty_unknown.sort(
        key=lambda item: (
            str(
                item["asset"].asset_name or ""
            ).lower()
        )
    )

    # ========================================================
    # HITUNG SUB-KATEGORI EXPIRING
    # ========================================================

    warranty_urgent = [
        item
        for item in warranty_expiring
        if (
            item["days_left"] is not None
            and item["days_left"] <= 7
        )
    ]

    warranty_warning = [
        item
        for item in warranty_expiring
        if (
            item["days_left"] is not None
            and item["days_left"] > 7
            and item["days_left"] <= 30
        )
    ]

    # ========================================================
    # PILIH DATA SESUAI FILTER
    # ========================================================

    if selected_status == "urgent":

        warranty_rows = warranty_urgent

    elif selected_status == "warning":

        warranty_rows = warranty_warning

    elif selected_status == "expiring":

        warranty_rows = warranty_expiring

    elif selected_status == "expired":

        warranty_rows = warranty_expired

    elif selected_status == "active":

        warranty_rows = warranty_active

    elif selected_status == "unknown":

        warranty_rows = warranty_unknown

    else:

        warranty_rows = (
            warranty_expiring
            + warranty_expired
            + warranty_active
            + warranty_unknown
        )

    # ========================================================
    # SORT DATA YANG DITAMPILKAN
    # ========================================================

    if selected_status in (
        "urgent",
        "warning",
        "expiring",
        "active",
    ):

        warranty_rows.sort(
            key=lambda item: (
                item["days_left"]
                if item["days_left"] is not None
                else 999999
            )
        )

    elif selected_status == "expired":

        warranty_rows.sort(
            key=lambda item: (
                item["end_date"]
                if item["end_date"] is not None
                else date.min
            ),
            reverse=True
        )

    # ========================================================
    # COUNT
    # ========================================================

    warranty_active_count = len(
        warranty_active
    )

    warranty_expiring_count = len(
        warranty_expiring
    )

    warranty_expired_count = len(
        warranty_expired
    )

    warranty_unknown_count = len(
        warranty_unknown
    )

    warranty_urgent_count = len(
        warranty_urgent
    )

    warranty_warning_count = len(
        warranty_warning
    )

    warranty_total_count = (
        warranty_active_count
        + warranty_expiring_count
        + warranty_expired_count
        + warranty_unknown_count
    )

    # ========================================================
    # RENDER
    # ========================================================

    return render_template(
        "warranty/alert.html",

        # ----------------------------------------------------
        # DATA UTAMA
        # ----------------------------------------------------

        warranty_rows=warranty_rows,

        warranty_active=warranty_active,

        warranty_expiring=warranty_expiring,

        warranty_expired=warranty_expired,

        warranty_unknown=warranty_unknown,

        # ----------------------------------------------------
        # SUB CATEGORY
        # ----------------------------------------------------

        warranty_urgent=warranty_urgent,

        warranty_warning=warranty_warning,

        # ----------------------------------------------------
        # COUNT
        # ----------------------------------------------------

        warranty_total_count=(
            warranty_total_count
        ),

        warranty_active_count=(
            warranty_active_count
        ),

        warranty_expiring_count=(
            warranty_expiring_count
        ),

        warranty_expired_count=(
            warranty_expired_count
        ),

        warranty_unknown_count=(
            warranty_unknown_count
        ),

        warranty_urgent_count=(
            warranty_urgent_count
        ),

        warranty_warning_count=(
            warranty_warning_count
        ),

        # ----------------------------------------------------
        # FILTER
        # ----------------------------------------------------

        selected_status=(
            selected_status
        ),

        search=search,

        today=today,
    )

# ============================================================
# WARRANTY NOTIFICATION CENTER
# ============================================================


@main.route("/warranty/notifications")
@requires_permission("warranty", "read")
def warranty_notifications():

    today = date.today()

    # ========================================================
    # FILTER
    # ========================================================

    selected_status = (
        request.args.get(
            "status",
            "all",
        )
        .strip()
        .lower()
    )

    search = (
        request.args.get(
            "search",
            "",
        )
        .strip()
    )

    # ========================================================
    # DEFAULT DATA
    # ========================================================

    urgent = []

    warning = []

    expired = []

    unknown = []

    # ========================================================
    # AMBIL SEMUA ASSET
    # ========================================================

    assets = (
        Asset.query
        .order_by(
            Asset.asset_name.asc(),
            Asset.id.asc(),
        )
        .all()
    )

    # ========================================================
    # PROSES WARRANTY
    # ========================================================

    for asset in assets:

        # ----------------------------------------------------
        # SEARCH
        # ----------------------------------------------------

        if search:

            searchable_text = " ".join([
                str(asset.asset_tag or ""),
                str(asset.asset_name or ""),
                str(asset.brand or ""),
                str(asset.model or ""),
                str(asset.serial_number or ""),
                str(asset.vendor or ""),
                str(asset.location or ""),
                str(asset.department or ""),
                str(asset.pic or ""),
            ]).lower()

            if search.lower() not in searchable_text:

                continue

        # ----------------------------------------------------
        # WARRANTY STATUS
        # ----------------------------------------------------

        warranty_info = get_warranty_status(
            asset.purchase_date,
            asset.warranty,
            today=today,
        )

        status = warranty_info.get(
            "status",
            "unknown",
        )

        days_left = warranty_info.get(
            "days_left",
        )

        end_date = warranty_info.get(
            "end_date",
        )

        item = {
            "asset": asset,
            "status": status,
            "days_left": days_left,
            "end_date": end_date,
        }

        # ----------------------------------------------------
        # URGENT
        # ----------------------------------------------------

        if (
            status == "expiring"
            and days_left is not None
            and days_left <= 7
        ):

            if selected_status in (
                "all",
                "urgent",
            ):

                urgent.append(item)

        # ----------------------------------------------------
        # WARNING
        # ----------------------------------------------------

        elif (
            status == "expiring"
            and days_left is not None
            and days_left > 7
            and days_left <= 30
        ):

            if selected_status in (
                "all",
                "warning",
            ):

                warning.append(item)

        # ----------------------------------------------------
        # EXPIRED
        # ----------------------------------------------------

        elif status == "expired":

            if selected_status in (
                "all",
                "expired",
            ):

                expired.append(item)

        # ----------------------------------------------------
        # UNKNOWN
        # ----------------------------------------------------

        elif status == "unknown":

            if selected_status in (
                "all",
                "unknown",
            ):

                unknown.append(item)

    # ========================================================
    # SORT URGENT
    # ========================================================

    urgent.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 999999
        )
    )

    # ========================================================
    # SORT WARNING
    # ========================================================

    warning.sort(
        key=lambda item: (
            item["days_left"]
            if item["days_left"] is not None
            else 999999
        )
    )

    # ========================================================
    # SORT EXPIRED
    # ========================================================

    expired.sort(
        key=lambda item: (
            item["end_date"]
            if item["end_date"] is not None
            else date.min
        ),
        reverse=True,
    )

    # ========================================================
    # SORT UNKNOWN
    # ========================================================

    unknown.sort(
        key=lambda item: (
            str(
                item["asset"].asset_name
                or ""
            ).lower()
        )
    )

    # ========================================================
    # COUNT HASIL FILTER
    # ========================================================

    urgent_count = len(urgent)

    warning_count = len(warning)

    expired_count = len(expired)

    unknown_count = len(unknown)

    total_alerts = (
        urgent_count
        + warning_count
        + expired_count
    )

    # ========================================================
    # RENDER
    # ========================================================

    return render_template(
        "warranty/notifications.html",

        urgent=urgent,

        warning=warning,

        expired=expired,

        unknown=unknown,

        urgent_count=urgent_count,

        warning_count=warning_count,

        expired_count=expired_count,

        unknown_count=unknown_count,

        total_alerts=total_alerts,

        selected_status=selected_status,

        search=search,

        today=today,
    )

# ============================================================
# WARRANTY HISTORY
# ============================================================


@main.route("/warranty/history/<int:asset_id>")
@requires_permission("warranty", "read")
def warranty_history(asset_id):

    asset = Asset.query.get_or_404(
        asset_id
    )

    histories = (
        WarrantyHistory.query
        .filter_by(
            asset_id=asset.id
        )
        .order_by(
            WarrantyHistory.changed_date.desc(),
            WarrantyHistory.id.desc()
        )
        .all()
    )

    return render_template(
        "warranty/history.html",
        asset=asset,
        histories=histories,
    )

# ============================================================
# EXPORT WARRANTY HISTORY
# ============================================================


@main.route("/warranty/history/<int:asset_id>/export")
@requires_permission("warranty", "read")
def export_warranty_history(asset_id):

    from openpyxl import Workbook
    from openpyxl.styles import (
        Font,
        Alignment,
        PatternFill,
    )

    asset = Asset.query.get_or_404(
        asset_id
    )

    histories = (
        WarrantyHistory.query
        .filter_by(
            asset_id=asset.id
        )
        .order_by(
            WarrantyHistory.changed_date.desc(),
            WarrantyHistory.id.desc()
        )
        .all()
    )

    # ========================================================
    # WORKBOOK
    # ========================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Warranty History"

    # ========================================================
    # TITLE
    # ========================================================

    worksheet.merge_cells("A1:H1")

    worksheet["A1"] = (
        "WARRANTY HISTORY - ASSET TRACKER"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=14,
    )

    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 26

    # ========================================================
    # ASSET INFORMATION
    # ========================================================

    worksheet.merge_cells("A2:H2")

    worksheet["A2"] = (
        f"Asset Tag: {asset.asset_tag or '-'}"
        f" | Asset Name: {asset.asset_name or '-'}"
    )

    worksheet["A2"].font = Font(
        bold=True,
    )

    worksheet["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.merge_cells("A3:H3")

    worksheet["A3"] = (
        f"Brand: {asset.brand or '-'}"
        f" | Model: {asset.model or '-'}"
        f" | Serial Number: {asset.serial_number or '-'}"
    )

    worksheet["A3"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # ========================================================
    # TABLE HEADER
    # ========================================================

    headers = [
        "No",
        "Tanggal Perubahan",
        "Warranty Lama",
        "Berakhir Lama",
        "Warranty Baru",
        "Berakhir Baru",
        "Alasan",
        "Catatan",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=5,
            column=column_number,
        )

        cell.value = header

        cell.font = Font(
            bold=True,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

    # ========================================================
    # DATA
    # ========================================================

    for number, history in enumerate(
        histories,
        start=1,
    ):

        worksheet.append([

            number,

            (
                history.changed_date
                if history.changed_date
                else ""
            ),

            history.old_warranty or "",

            (
                history.old_end_date
                if history.old_end_date
                else ""
            ),

            history.new_warranty or "",

            (
                history.new_end_date
                if history.new_end_date
                else ""
            ),

            history.reason or "",

            history.notes or "",

        ])

    # ========================================================
    # FORMAT DATE
    # ========================================================

    for row in worksheet.iter_rows(
        min_row=6,
        max_row=worksheet.max_row,
    ):

        # B = Tanggal perubahan
        if row[1].value:

            row[1].number_format = (
                "dd-mm-yyyy"
            )

        # D = End date lama
        if row[3].value:

            row[3].number_format = (
                "dd-mm-yyyy"
            )

        # F = End date baru
        if row[5].value:

            row[5].number_format = (
                "dd-mm-yyyy"
            )

    # ========================================================
    # ALIGNMENT
    # ========================================================

    for row in worksheet.iter_rows(
        min_row=5,
        max_row=worksheet.max_row,
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="center",
            )

    for row in worksheet.iter_rows(
        min_row=6,
        max_row=worksheet.max_row,
    ):

        for index in (
            0,
            1,
            3,
            5,
        ):

            row[index].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # ========================================================
    # FREEZE HEADER
    # ========================================================

    worksheet.freeze_panes = "A6"

    # ========================================================
    # AUTO FILTER
    # ========================================================

    if worksheet.max_row >= 5:

        worksheet.auto_filter.ref = (
            f"A5:H{worksheet.max_row}"
        )

    # ========================================================
    # COLUMN WIDTH
    # ========================================================

    column_widths = {

        "A": 8,
        "B": 20,
        "C": 22,
        "D": 20,
        "E": 22,
        "F": 20,
        "G": 32,
        "H": 45,

    }

    for column, width in column_widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # ========================================================
    # OUTPUT
    # ========================================================

    output = io.BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    filename = (
        "warranty_history_"
        + str(asset.asset_tag or asset.id)
        + "_"
        + datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )
        + ".xlsx"
    )

    return send_file(

        output,

        as_attachment=True,

        download_name=filename,

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

    )

# ============================================================
# ASSET LIST
# ============================================================


